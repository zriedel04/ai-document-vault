"""The DOCX extractor, on synthetic documents built here with python-docx and on hand-built zips."""

import io
import struct
import tracemalloc
import zipfile
import zlib
from collections.abc import Callable
from typing import Any, NoReturn

import docx
import pytest
from docx.document import Document as DocumentObject
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from vault.ingest.extractors import Extraction, ExtractionError
from vault.ingest.extractors import docx as extractor
from vault.ingest.extractors.docx import (
    CELL_SEPARATOR,
    CHARACTERS,
    EXCLUDED,
    PARAGRAPH_BREAK,
    VERSION,
    CompatibilityError,
    NotDocxError,
    extract,
)
from vault.ingest.validate import OLE_SIGNATURE, Limits, Reason, check_docx

# Letters outside ASCII, from code points: e acute, an em dash, the CJK word for resume, and a
# character outside the Basic Multilingual Plane (one Python character, two UTF-16 units).
WORDS = "café — 履歴書 \U0002000b"
CANARY = "canary-secret-xyzzy"
W = nsdecls("w")
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
WPS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
VML = 'xmlns:v="urn:schemas-microsoft-com:vml"'
MATH = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
FUTURE = 'xmlns:x="urn:example:future-wrapper"'


def build(fill: Callable[[DocumentObject], object] | None = None) -> bytes:
    document = docx.Document()
    if fill is not None:
        fill(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def text_of(data: bytes) -> str:
    result = extract(data)
    assert len(result.pages) == 1
    return result.pages[0].text


def raw_zip(members: list[tuple[str, bytes, int, int, int]]) -> bytes:
    """A zip written by hand: (name, stored bytes, method, CRC, claimed size) per member, so the
    claims can differ from the data."""
    body, directory = b"", b""
    for name, stored, method, crc, size in members:
        encoded = name.encode()
        offset = len(body)
        header = struct.pack(
            "<4s5H3L2H",
            b"PK\x03\x04",
            20,
            0,
            method,
            0,
            33,
            crc,
            len(stored),
            size,
            len(encoded),
            0,
        )
        body += header + encoded + stored
        directory += (
            struct.pack(
                "<4s6H3L5H2L",
                b"PK\x01\x02",
                20,
                20,
                0,
                method,
                0,
                33,
                crc,
                len(stored),
                size,
                len(encoded),
                0,
                0,
                0,
                0,
                0,
                offset,
            )
            + encoded
        )
    count = len(members)
    end = struct.pack("<4s4H2LH", b"PK\x05\x06", 0, 0, count, count, len(directory), len(body), 0)
    return body + directory + end


def members_of(data: bytes) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return [(info.filename, archive.read(info)) for info in archive.infolist()]


def deflate(data: bytes) -> bytes:
    compressor = zlib.compressobj(9, zlib.DEFLATED, -15)
    return compressor.compress(data) + compressor.flush()


def replace_member(data: bytes, name: str, content: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for member, original in members_of(data):
            archive.writestr(member, content if member == name else original)
    return out.getvalue()


def add_xml(paragraph_owner: DocumentObject, xml: str) -> None:
    """Append a paragraph written as XML to the body, before the section properties."""
    body = paragraph_owner.element.body
    body.insert(len(body) - 1, parse_xml(xml))


@pytest.fixture
def library_never_called(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(*_: object, **__: object) -> NoReturn:
        raise AssertionError("python-docx was called")

    monkeypatch.setattr(extractor, "Document", refuse)
    monkeypatch.setattr(extractor, "_bounded_copy", refuse)


# The result


def test_paragraphs_in_order_one_page_and_the_version() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_paragraph("First paragraph.")
        document.add_paragraph("")  # empty paragraphs are left out
        document.add_paragraph("   ")
        document.add_paragraph("Second paragraph.")

    result = extract(build(fill))

    assert isinstance(result, Extraction)
    assert result.extractor == VERSION == "docx-1"
    assert result.encoding is None
    assert [page.number for page in result.pages] == [1]
    assert result.pages[0].text == f"First paragraph.{PARAGRAPH_BREAK}Second paragraph."


def test_an_empty_document_is_one_empty_page() -> None:
    result = extract(build())

    assert result.pages[0].text == ""
    assert result.pages[0].headings == ()


def test_unicode_is_kept_and_offsets_count_characters() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_paragraph(WORDS)
        document.add_heading(WORDS, level=2)

    page = extract(build(fill)).pages[0]

    assert page.text == f"{WORDS}{PARAGRAPH_BREAK}{WORDS}"
    (heading,) = page.headings
    assert page.text[heading.start : heading.end] == WORDS
    assert heading.start == len(WORDS) + len(PARAGRAPH_BREAK)


def test_line_breaks_inside_a_paragraph_are_newlines() -> None:
    def fill(document: DocumentObject) -> None:
        paragraph = document.add_paragraph("one")
        paragraph.add_run().add_break()
        paragraph.add_run("two\rthree")  # python-docx writes \r as a break as well
        add_xml(
            document, f"<w:p {W}><w:r><w:t>four</w:t><w:cr/><w:t>five&#13;six</w:t></w:r></w:p>"
        )

    assert text_of(build(fill)) == f"one\ntwo\nthree{PARAGRAPH_BREAK}four\nfive\nsix"


def test_tabs_and_special_characters_in_runs() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W}><w:r><w:t>a</w:t><w:tab/><w:t>b</w:t><w:noBreakHyphen/>"
            "<w:t>c</w:t><w:softHyphen/><w:t>d</w:t></w:r></w:p>",
        )

    assert text_of(build(fill)) == "a\tb-cd"


# Headings


def test_headings_at_every_level_and_the_title() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_heading("The title", level=0)
        for level in range(1, 10):
            document.add_heading(f"Level {level}", level=level)
            document.add_paragraph(f"Body under {level}.")

    page = extract(build(fill)).pages[0]

    assert [(h.level, h.title) for h in page.headings] == [
        (1, "The title"),
        *((level, f"Level {level}") for level in range(1, 10)),
    ]


def test_every_heading_offset_points_at_its_heading_text() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_paragraph("Before any heading.")
        document.add_heading(f"Intro {WORDS}", level=1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "x"
        document.add_heading("  Spaced   title  ", level=2)
        heading = document.add_heading("Broken", level=3)
        heading.add_run().add_break()
        heading.add_run("over lines")
        document.add_paragraph("")
        document.add_heading("Last", level=1)

    page = extract(build(fill)).pages[0]

    assert [h.title for h in page.headings] == [
        f"Intro {WORDS}",
        "Spaced title",
        "Broken over lines",
        "Last",
    ]
    for heading in page.headings:
        written = page.text[heading.start : heading.end]
        assert " ".join(written.split()) == heading.title
        assert heading.start == 0 or page.text[heading.start - 2 : heading.start] == "\n\n"
        assert heading.end == len(page.text) or page.text[heading.end : heading.end + 2] == "\n\n"
    assert [h.start for h in page.headings] == sorted(h.start for h in page.headings)


def test_only_the_built_in_heading_styles_are_headings() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_paragraph("Subtitle text", style="Subtitle")
        document.add_paragraph("Looks like a heading", style="Intense Quote")
        run_styled = document.add_paragraph().add_run("character style")
        run_styled.style = document.styles["Heading 1 Char"]
        add_xml(
            document,
            f'<w:p {W}><w:pPr><w:pStyle w:val="NoSuchStyle"/></w:pPr>'
            "<w:r><w:t>unknown style</w:t></w:r></w:p>",
        )

    assert extract(build(fill)).pages[0].headings == ()


def test_a_heading_style_is_found_by_name_not_by_id() -> None:
    # A Word in another language writes another style id for the same built-in name.
    def fill(document: DocumentObject) -> None:
        document.add_heading("Kapitel", level=2)

    data = build(fill)
    styles = dict(members_of(data))["word/styles.xml"]
    document_xml = dict(members_of(data))["word/document.xml"]
    renamed = replace_member(
        data, "word/styles.xml", styles.replace(b'"Heading2"', b'"berschrift2"')
    )
    renamed = replace_member(
        renamed, "word/document.xml", document_xml.replace(b'"Heading2"', b'"berschrift2"')
    )

    assert [(h.level, h.title) for h in extract(renamed).pages[0].headings] == [(2, "Kapitel")]


def test_a_style_that_is_not_a_paragraph_style_is_not_a_heading() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_heading("Not a section", level=2)

    data = build(fill)
    styles = dict(members_of(data))["word/styles.xml"]
    old = b'<w:style w:type="paragraph" w:styleId="Heading2">'
    assert old in styles
    changed = styles.replace(old, b'<w:style w:type="character" w:styleId="Heading2">')

    assert extract(replace_member(data, "word/styles.xml", changed)).pages[0].headings == ()


def test_a_heading_inside_a_table_or_a_text_box_is_not_a_section_heading() -> None:
    def fill(document: DocumentObject) -> None:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).paragraphs[0].text = "In a cell"
        table.cell(0, 0).paragraphs[0].style = document.styles["Heading 1"]
        add_xml(document, text_box_paragraph("Boxed heading", style="Heading1"))

    page = extract(build(fill)).pages[0]

    assert "In a cell" in page.text
    assert "Boxed heading" in page.text
    assert page.headings == ()


# Tables


def test_a_table_is_one_line_per_row_with_tab_separated_cells() -> None:
    def fill(document: DocumentObject) -> None:
        document.add_paragraph("Before.")
        table = document.add_table(rows=3, cols=3)
        for row, values in enumerate([("Name", "Hours", "Note"), ("Alex", "41", ""), ("", "", "")]):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        document.add_paragraph("After.")

    assert CELL_SEPARATOR == "\t"
    assert text_of(build(fill)) == (
        f"Before.{PARAGRAPH_BREAK}Name\tHours\tNote\nAlex\t41\t{PARAGRAPH_BREAK}After."
    )


def test_cell_paragraphs_breaks_tabs_and_nested_tables_stay_on_the_row() -> None:
    def fill(document: DocumentObject) -> None:
        table = document.add_table(rows=1, cols=2)
        cell = table.cell(0, 0)
        cell.text = "first"
        cell.add_paragraph("second\tpart")
        cell.paragraphs[1].add_run().add_break()
        cell.paragraphs[1].add_run("third")
        inner = table.cell(0, 1).add_table(rows=2, cols=2)
        for row in range(2):
            for column in range(2):
                inner.cell(row, column).text = f"r{row}c{column}"

    assert text_of(build(fill)) == "first second part third\tr0c0 r0c1 r1c0 r1c1"


def test_a_merged_cell_appears_once() -> None:
    def fill(document: DocumentObject) -> None:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "wide"
        table.cell(1, 0).text = "a"
        table.cell(1, 1).text = "b"

    assert text_of(build(fill)) == "wide\na\tb"


# Tracked changes


def test_tracked_changes_are_read_as_accepted() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W}><w:r><w:t xml:space="preserve">Keep </w:t></w:r>'
            '<w:ins w:id="1" w:author="A"><w:r><w:t>inserted</w:t></w:r></w:ins>'
            # A deleted run keeps its tabs and breaks as w:tab and w:br, not as deleted text.
            f'<w:del w:id="2" w:author="A"><w:r><w:delText>{CANARY}</w:delText><w:tab/><w:br/>'
            "</w:r></w:del>"
            f'<w:moveFrom w:id="3" w:author="A"><w:r><w:t>{CANARY}-moved</w:t></w:r></w:moveFrom>'
            '<w:moveTo w:id="4" w:author="A"><w:r><w:t xml:space="preserve"> moved in</w:t>'
            "</w:r></w:moveTo></w:p>",
        )

    text = text_of(build(fill))

    assert text == "Keep inserted moved in"
    assert CANARY not in text


def test_a_row_deleted_as_a_tracked_change_is_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "kept row"
        table.cell(1, 0).text = CANARY
        row = table.rows[1]._tr
        row.insert(0, parse_xml(f'<w:trPr {W}><w:del w:id="9" w:author="A"/></w:trPr>'))

    assert text_of(build(fill)) == "kept row"


def test_a_cell_deleted_as_a_tracked_change_is_left_out() -> None:
    """Accepting a cell deletion removes the cell: its text goes and the next cell moves up. A
    row left with only deleted cells goes too, an inserted cell stays, and a nested table follows
    the same rule."""

    def mark(cell: Any, revision: str) -> None:
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:{revision} {W} w:id="7" w:author="A"/>'))

    def fill(document: DocumentObject) -> None:
        table = document.add_table(rows=3, cols=3)
        for row, values in enumerate(
            (("a", CANARY, "c"), (CANARY, CANARY, CANARY), ("x", "y", "z"))
        ):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        mark(table.cell(0, 1), "cellDel")
        for column in range(3):
            mark(table.cell(1, column), "cellDel")
        mark(table.cell(2, 2), "cellIns")
        nested = table.cell(2, 0).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = CANARY
        nested.cell(0, 1).text = "inner"
        mark(nested.cell(0, 0), "cellDel")

    text = text_of(build(fill))
    assert text == "a\tc\nx inner\ty\tz"
    assert CANARY not in text


# What is left out


def test_hidden_text_is_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        paragraph = document.add_paragraph("Visible ")
        paragraph.add_run(CANARY).font.hidden = True
        paragraph.add_run("text")
        add_xml(
            document,
            f'<w:p {W}><w:r><w:rPr><w:vanish w:val="false"/></w:rPr>'
            "<w:t>shown again</w:t></w:r></w:p>",
        )

    assert text_of(build(fill)) == f"Visible text{PARAGRAPH_BREAK}shown again"


def test_headers_and_footers_are_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        document.sections[0].header.paragraphs[0].text = f"{CANARY} header"
        document.sections[0].footer.paragraphs[0].text = f"{CANARY} footer"
        document.add_paragraph("Body.")

    assert text_of(build(fill)) == "Body."


def test_comments_are_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        paragraph = document.add_paragraph("Commented text.")
        document.add_comment(paragraph.runs, text=f"{CANARY} comment", author="A")

    data = build(fill)

    assert "word/comments.xml" in dict(members_of(data))
    assert text_of(data) == "Commented text."


def test_footnotes_are_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        notes = (
            f'<w:footnotes {W}><w:footnote w:id="1"><w:p><w:r><w:t>{CANARY} note</w:t></w:r>'
            "</w:p></w:footnote></w:footnotes>"
        ).encode()
        part = Part(
            PackURI("/word/footnotes.xml"),
            CONTENT_TYPE.WML_FOOTNOTES,
            notes,
            document.part.package,
        )
        document.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
        add_xml(
            document,
            f'<w:p {W}><w:r><w:t>Claim</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>',
        )

    data = build(fill)

    assert "word/footnotes.xml" in dict(members_of(data))
    assert text_of(data) == "Claim"


def test_field_codes_are_left_out_and_results_and_links_kept() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            f'<w:r><w:instrText xml:space="preserve"> HYPERLINK "{CANARY}" </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            "<w:r><w:t>field result</w:t></w:r>"
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
            '<w:r><w:t xml:space="preserve">, </w:t></w:r>'
            '<w:hyperlink r:id="rIdMissing"><w:r><w:t>link text</w:t></w:r></w:hyperlink>'
            '<w:fldSimple w:instr=" PAGE "><w:r><w:t>, 7</w:t></w:r></w:fldSimple></w:p>',
        )

    assert text_of(build(fill)) == "field result, link text, 7"


def test_content_controls_are_read_through() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:sdt {W}><w:sdtPr><w:alias w:val="{CANARY}"/></w:sdtPr><w:sdtContent>'
            "<w:p><w:r><w:t>block control</w:t></w:r>"
            '<w:sdt><w:sdtContent><w:r><w:t xml:space="preserve"> and inline</w:t></w:r>'
            "</w:sdtContent></w:sdt></w:p></w:sdtContent></w:sdt>",
        )

    assert text_of(build(fill)) == "block control and inline"


def test_custom_xml_wrappers_are_read_through() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:customXml {W} w:element="clause"><w:p><w:r><w:t>wrapped</w:t></w:r></w:p>'
            "</w:customXml>",
        )

    assert text_of(build(fill)) == "wrapped"


def test_alternate_content_in_a_paragraph_is_read_from_the_first_choice_only() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W} {MC}><mc:AlternateContent><mc:Choice Requires="w">'
            f"<w:r><w:t>new form</w:t></w:r></mc:Choice><mc:Fallback><w:r><w:t>{CANARY}</w:t>"
            "</w:r></mc:Fallback></mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)) == "new form"


def test_alternate_content_with_only_a_fallback_reads_the_fallback() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {MC}><mc:AlternateContent><mc:Fallback><w:r><w:t>only form</w:t></w:r>"
            "</mc:Fallback></mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)) == "only form"


W14 = 'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
W15 = 'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"'


@pytest.mark.parametrize(
    ("choices", "expected"),
    [
        pytest.param(
            [("w15", CANARY), ("wps", "understood")], "understood", id="next choice understood"
        ),
        pytest.param([("w15", CANARY)], "fallback", id="none understood, fallback read"),
        pytest.param([("w wps", "both known")], "both known", id="every prefix understood"),
        pytest.param([("w w15", CANARY)], "fallback", id="one prefix not understood"),
        pytest.param([("x", CANARY)], "fallback", id="namespace of another kind"),
        pytest.param(
            [("wps", "taken"), ("nowhere", CANARY)], "taken", id="choice after the taken one unread"
        ),
    ],
)
def test_a_choice_is_read_only_when_every_namespace_it_requires_is_understood(
    choices: list[tuple[str | None, str]], expected: str
) -> None:
    def fill(document: DocumentObject) -> None:
        options = "".join(
            f"<mc:Choice{'' if requires is None else ' Requires=' + repr(requires)}>{run(text)}"
            "</mc:Choice>"
            for requires, text in choices
        )
        add_xml(
            document,
            f"<w:p {W} {MC} {WPS} {W15} {FUTURE}><mc:AlternateContent>{options}"
            f"<mc:Fallback>{run('fallback')}</mc:Fallback></mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)) == expected


@pytest.mark.parametrize(
    "requires",
    [
        pytest.param(' Requires="nowhere"', id="prefix not in scope"),
        pytest.param(' Requires="w nowhere"', id="one prefix not in scope"),
        pytest.param(' Requires=""', id="nothing required"),
        pytest.param(' Requires=" "', id="only white space"),
        pytest.param("", id="no Requires"),
    ],
)
def test_a_choice_whose_requires_is_an_error_refuses_the_file(requires: str) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {MC}><mc:AlternateContent><mc:Choice{requires}>{run(CANARY)}</mc:Choice>"
            f"<mc:Fallback>{run('fallback')}</mc:Fallback></mc:AlternateContent></w:p>",
        )

    with pytest.raises(CompatibilityError):
        extract(build(fill))


def test_alternate_content_with_no_readable_choice_and_no_fallback_gives_nothing() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W} {MC} {W15}>{run("[")}<mc:AlternateContent><mc:Choice Requires="w15">'
            f"{run(CANARY)}</mc:Choice></mc:AlternateContent>{run(']')}</w:p>",
        )

    assert text_of(build(fill)) == "[]"


ALT_CHUNK = '<w:altChunk r:id="rIdImported"/>'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


@pytest.mark.parametrize(
    "block",
    [
        pytest.param(f"<w:body {W} {R}>{ALT_CHUNK}</w:body>", id="body"),
        pytest.param(
            f"<w:body {W} {R}><w:tbl><w:tr><w:tc>{ALT_CHUNK}<w:p/></w:tc></w:tr></w:tbl></w:body>",
            id="table cell",
        ),
        pytest.param(
            f"<w:body {W} {R}><w:sdt><w:sdtContent>{ALT_CHUNK}</w:sdtContent></w:sdt></w:body>",
            id="content control",
        ),
        pytest.param(
            f'<w:body {W} {R}><w:customXml w:element="x">{ALT_CHUNK}</w:customXml></w:body>',
            id="custom xml",
        ),
        pytest.param(
            f'<w:body {W} {R} {MC}><mc:AlternateContent><mc:Choice Requires="w">{ALT_CHUNK}'
            "</mc:Choice></mc:AlternateContent></w:body>",
            id="alternate content",
        ),
        pytest.param(
            f'<w:body {W} {R} {MC} {WPS}><w:p><w:r><mc:AlternateContent><mc:Choice Requires="wps">'
            f"<w:drawing><wps:wsp><wps:txbx><w:txbxContent>{ALT_CHUNK}</w:txbxContent></wps:txbx>"
            "</wps:wsp></w:drawing></mc:Choice></mc:AlternateContent></w:r></w:p></w:body>",
            id="text box",
        ),
        pytest.param(f"<w:body {W} {R}><w:p>{ALT_CHUNK}</w:p></w:body>", id="inside a paragraph"),
    ],
)
def test_content_imported_from_another_file_is_refused_wherever_it_is(block: str) -> None:
    def fill(document: DocumentObject) -> None:
        body = document.element.body
        add_xml(document, f"<w:p {W}>{run(CANARY)}</w:p>")
        for element in parse_xml(block):
            body.insert(len(body) - 1, element)

    with pytest.raises(extractor.ImportedContentError) as raised:
        extract(build(fill))
    assert isinstance(raised.value, ExtractionError)
    assert CANARY not in str(raised.value)
    assert "rIdImported" not in str(raised.value)


def test_imported_content_in_a_representation_that_is_not_read_is_not_refused() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W} {MC} {R}><mc:AlternateContent><mc:Choice Requires="w">{run("shown")}'
            f"</mc:Choice><mc:Fallback>{ALT_CHUNK}</mc:Fallback></mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)) == "shown"


def alternate(choice: str, fallback: str = "") -> str:
    tail = f"<mc:Fallback>{fallback}</mc:Fallback>" if fallback else ""
    head = '<mc:AlternateContent><mc:Choice Requires="w">'
    return f"{head}{choice}</mc:Choice>{tail}</mc:AlternateContent>"


def cell(text: str) -> str:
    return f"<w:tc><w:p>{run(text)}</w:p></w:tc>"


def test_alternate_content_around_blocks_rows_and_cells_is_read_once() -> None:
    """Alternate content may wrap a paragraph or a table in the body, a row in a table and a cell
    in a row: the choice is read, as in a paragraph, and the fallback copy is not."""
    fallback_paragraph = f"<w:p>{run(CANARY)}</w:p>"

    def fill(document: DocumentObject) -> None:
        body = document.element.body
        blocks = [
            f"<w:p>{run('paragraph')}</w:p>",
            f"<w:tbl><w:tr>{cell('table')}</w:tr></w:tbl>",
        ]
        for block in blocks:
            wrapped = alternate(block, fallback_paragraph)
            body.insert(len(body) - 1, parse_xml(f"<w:body {W} {MC}>{wrapped}</w:body>")[0])
        row = alternate(f"<w:tr>{cell('row')}</w:tr>", f"<w:tr>{cell(CANARY)}</w:tr>")
        cells = f"<w:tr>{cell('a')}{alternate(cell('b'), cell(CANARY))}</w:tr>"
        table = f"<w:tbl {W} {MC}>{row}{cells}</w:tbl>"
        body.insert(len(body) - 1, parse_xml(table))

    text = text_of(build(fill))
    assert text == PARAGRAPH_BREAK.join(["paragraph", "table", "row\na\tb"])
    assert CANARY not in text


# Every container of run text, decided one by one


def run(text: str) -> str:
    return f"<w:r><w:t>{text}</w:t></w:r>"


def ruby(base: str, guide: str) -> str:
    return (
        f'<w:r><w:ruby><w:rubyPr><w:rubyAlign w:val="center"/></w:rubyPr>'
        f"<w:rt><w:r><w:t>{guide}</w:t></w:r></w:rt>"
        f"<w:rubyBase><w:r><w:t>{base}</w:t></w:r></w:rubyBase></w:ruby></w:r>"
    )


# (what it is, the paragraph's content, the text it must give). CANARY marks text left out on
# purpose; every other word must come through.
INCLUDED_CONTAINERS = [
    ("ruby base", ruby("\u6f22\u5b57", f"{CANARY}-kana"), "\u6f22\u5b57"),
    ("hyperlink", f"<w:hyperlink>{run('linked')}</w:hyperlink>", "linked"),
    (
        "smart tag",
        f'<w:smartTag w:uri="urn:x" w:element="place">{run("Tempe")}</w:smartTag>',
        "Tempe",
    ),
    ("inline custom XML", f'<w:customXml w:element="x">{run("tagged")}</w:customXml>', "tagged"),
    (
        "inline content control",
        f'<w:sdt><w:sdtPr><w:alias w:val="{CANARY}"/></w:sdtPr><w:sdtContent>{run("ctrl")}'
        "</w:sdtContent></w:sdt>",
        "ctrl",
    ),
    ("simple field", f'<w:fldSimple w:instr=" PAGE {CANARY} ">{run("7")}</w:fldSimple>', "7"),
    ("insertion", f'<w:ins w:id="1" w:author="A">{run("added")}</w:ins>', "added"),
    ("move, new place", f'<w:moveTo w:id="2" w:author="A">{run("moved")}</w:moveTo>', "moved"),
    (
        "conflict insertion",
        f'<w14:conflictIns w:id="6" w:author="A">{run("agreed")}</w14:conflictIns>',
        "agreed",
    ),
    ("text direction", f'<w:dir w:val="rtl">{run("rtl")}</w:dir>', "rtl"),
    ("bidirectional override", f'<w:bdo w:val="ltr">{run("ltr")}</w:bdo>', "ltr"),
    (
        "alternate content in a paragraph",
        f'<mc:AlternateContent><mc:Choice Requires="w">{run("choice")}</mc:Choice>'
        f"<mc:Fallback>{run(CANARY)}</mc:Fallback></mc:AlternateContent>",
        "choice",
    ),
    (
        "alternate content in a run",
        f'<w:r><mc:AlternateContent><mc:Choice Requires="w"><w:t>in run</w:t></mc:Choice>'
        f"<mc:Fallback><w:t>{CANARY}</w:t></mc:Fallback></mc:AlternateContent></w:r>",
        "in run",
    ),
    ("a wrapper this module does not know", f"<x:later>{run('kept')}</x:later>", "kept"),
    (
        "characters",
        '<w:r><w:t>a</w:t><w:tab/><w:t>b</w:t><w:ptab w:relativeTo="margin" w:alignment="right"'
        ' w:leader="none"/><w:t>c</w:t><w:br/><w:t>d</w:t><w:cr/><w:t>e</w:t><w:noBreakHyphen/>'
        "<w:t>f</w:t><w:softHyphen/><w:t>g</w:t></w:r>",
        "a\tb\tc\nd\ne-fg",
    ),
    (
        "markers with no text",
        '<w:bookmarkStart w:id="0" w:name="m"/><w:proofErr w:type="spellStart"/>'
        f'{run("marked")}<w:commentRangeStart w:id="5"/><w:bookmarkEnd w:id="0"/>'
        '<w:r><w:footnoteReference w:id="1"/><w:lastRenderedPageBreak/></w:r>',
        "marked",
    ),
]
EXCLUDED_CONTAINERS = [
    ("tracked deletion", f'<w:del w:id="3" w:author="A">{run(CANARY)}</w:del>', "w:del"),
    ("deleted text", f"<w:r><w:delText>{CANARY}</w:delText></w:r>", "w:delText"),
    (
        "move, old place",
        f'<w:moveFrom w:id="4" w:author="A">{run(CANARY)}</w:moveFrom>',
        "w:moveFrom",
    ),
    (
        "conflict deletion",
        f'<w14:conflictDel w:id="7" w:author="A">{run(CANARY)}'
        f"<w:r><w:delText>{CANARY}</w:delText></w:r></w14:conflictDel>",
        "w14:conflictDel",
    ),
    ("field code", f"<w:r><w:instrText>{CANARY}</w:instrText></w:r>", "w:instrText"),
    (
        "deleted field code",
        f"<w:r><w:delInstrText>{CANARY}</w:delInstrText></w:r>",
        "w:delInstrText",
    ),
    ("ruby guide", ruby("", CANARY), "w:rt"),
    (
        "content control properties",
        f"<w:sdt><w:sdtPr>{run(CANARY)}</w:sdtPr><w:sdtEndPr>{run(CANARY)}</w:sdtEndPr>"
        "<w:sdtContent/></w:sdt>",
        "w:sdtPr w:sdtEndPr",
    ),
    ("symbol", '<w:r><w:sym w:font="Wingdings" w:char="F0E0"/></w:r>', "w:sym"),
    (
        "equation",
        f"<m:oMathPara><m:oMath><m:r><m:t>x</m:t></m:r>{run(CANARY)}</m:oMath></m:oMathPara>",
        "m:oMathPara m:oMath",
    ),
    ("properties", f"<w:pPr><w:rPr>{run(CANARY)}</w:rPr></w:pPr>", "w:pPr w:rPr"),
]


def paragraph(content: str) -> str:
    return f"<w:p {W} {MC} {MATH} {W14} {FUTURE}>{content}</w:p>"


@pytest.mark.parametrize(
    ("content", "expected"),
    [pytest.param(content, expected, id=name) for name, content, expected in INCLUDED_CONTAINERS],
)
def test_text_inside_each_container_is_kept(content: str, expected: str) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, paragraph(f"{run('[')}{content}{run(']')}"))

    assert text_of(build(fill)) == f"[{expected}]"


@pytest.mark.parametrize(
    "content",
    [pytest.param(content, id=name) for name, content, _ in EXCLUDED_CONTAINERS],
)
def test_text_inside_each_excluded_element_is_left_out(content: str) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, paragraph(f"{run('[')}{content}{run(']')}"))

    assert text_of(build(fill)) == "[]"


def test_every_excluded_element_has_a_test() -> None:
    prefixes = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    }
    tested = {
        "{%s}%s" % (prefixes[name.split(":")[0]], name.split(":")[1])  # noqa: UP031
        for _, _, names in EXCLUDED_CONTAINERS
        for name in names.split()
    }
    assert tested == set(EXCLUDED)
    # The "characters" case above writes each of these.
    assert {tag.split("}")[1] for tag in CHARACTERS} == {"tab", "ptab", "br", "cr", "noBreakHyphen"}


# Every element ECMA-376 Part 1 (17.13.5) and Word 2010's extensions (w14) use for inline content
# that accepting the changes removes, and the elements for content it keeps. Range markers
# (``w:moveFromRangeStart`` and the like) hold no text, and properties are left out whole.
DELETIONS = {"w:del", "w:moveFrom", "w:delText", "w:delInstrText", "w14:conflictDel"}
INSERTIONS = {"w:ins", "w:moveTo", "w14:conflictIns"}


def test_every_kind_of_deleted_content_is_excluded_and_inserted_content_is_not() -> None:
    namespaces = {
        "w": "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}",
        "w14": "{http://schemas.microsoft.com/office/word/2010/wordml}",
    }

    def clark(name: str) -> str:
        prefix, local = name.split(":")
        return namespaces[prefix] + local

    assert {clark(name) for name in DELETIONS} <= EXCLUDED
    assert not {clark(name) for name in INSERTIONS} & EXCLUDED


def test_included_containers_nested_in_each_other_still_give_their_text() -> None:
    # Every pair of containers, one inside the other: no combination may lose the text.
    wrappers = [
        ("<w:hyperlink>", "</w:hyperlink>"),
        ('<w:smartTag w:element="x">', "</w:smartTag>"),
        ('<w:customXml w:element="x">', "</w:customXml>"),
        ("<w:sdt><w:sdtContent>", "</w:sdtContent></w:sdt>"),
        ('<w:fldSimple w:instr=" X ">', "</w:fldSimple>"),
        ('<w:ins w:id="1" w:author="A">', "</w:ins>"),
        ('<w:moveTo w:id="2" w:author="A">', "</w:moveTo>"),
        ('<w:dir w:val="rtl">', "</w:dir>"),
        ('<w:bdo w:val="ltr">', "</w:bdo>"),
        ('<mc:AlternateContent><mc:Choice Requires="w">', "</mc:Choice></mc:AlternateContent>"),
        ("<x:later>", "</x:later>"),
    ]
    pairs = [(outer, inner) for outer in wrappers for inner in wrappers]

    def fill(document: DocumentObject) -> None:
        for number, ((open_a, close_a), (open_b, close_b)) in enumerate(pairs):
            add_xml(document, paragraph(f"{open_a}{open_b}{run(f'w{number}')}{close_b}{close_a}"))

    text = text_of(build(fill))

    assert text.split(PARAGRAPH_BREAK) == [f"w{number}" for number in range(len(pairs))]


def test_a_field_nested_in_another_fields_code_is_left_out() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            paragraph(
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                '<w:r><w:instrText xml:space="preserve"> IF </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                "<w:r><w:instrText> DATE </w:instrText></w:r>"
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                f"{run(CANARY)}"
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                f"{run('shown')}"
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
                f"{run(' after')}"
            ),
        )

    assert text_of(build(fill)) == "shown after"


def test_a_field_left_open_ends_with_its_paragraph() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, paragraph(f'<w:r><w:fldChar w:fldCharType="begin"/></w:r>{run(CANARY)}'))
        add_xml(document, paragraph(f'<w:r><w:fldChar w:fldCharType="end"/></w:r>{run("next")}'))

    assert text_of(build(fill)) == "next"


# Field markers where the text around them is left out


def fld(kind: str) -> str:
    return f'<w:r><w:fldChar w:fldCharType="{kind}"/></w:r>'


def code(text: str) -> str:
    return f'<w:r><w:instrText xml:space="preserve"> {text} </w:instrText></w:r>'


# Places that are in the document as accepted but not shown: the text there is left out, and a
# field marker there still opens, splits or closes its field.
NOT_SHOWN: dict[str, Callable[[str], str]] = {
    "hidden run": lambda kind: (
        f'<w:r><w:rPr><w:vanish/></w:rPr><w:fldChar w:fldCharType="{kind}"/>'
        f"<w:t>{CANARY}</w:t></w:r>"
    ),
    "hidden run inside a hyperlink": lambda kind: (
        f'<w:hyperlink><w:r><w:rPr><w:vanish w:val="1"/></w:rPr>'
        f'<w:fldChar w:fldCharType="{kind}"/></w:r></w:hyperlink>'
    ),
    "ruby guide": lambda kind: (
        f"<w:r><w:ruby><w:rt>{fld(kind)}{run(CANARY)}</w:rt><w:rubyBase/></w:ruby></w:r>"
    ),
    "equation": lambda kind: (
        f"<m:oMathPara><m:oMath>{fld(kind)}{run(CANARY)}</m:oMath></m:oMathPara>"
    ),
}
# Places whose markers are not in the document as accepted and read: content accepting the
# changes removes, and a representation of alternate content that is not read.
NOT_PRESENT: dict[str, Callable[[str], str]] = {
    "tracked deletion": lambda kind: f'<w:del w:id="1" w:author="A">{fld(kind)}</w:del>',
    "move, old place": lambda kind: f'<w:moveFrom w:id="2" w:author="A">{fld(kind)}</w:moveFrom>',
    "conflict deletion": lambda kind: (
        f'<w14:conflictDel w:id="3" w:author="A">{fld(kind)}</w14:conflictDel>'
    ),
    "fallback not read": lambda kind: (
        f'<mc:AlternateContent><mc:Choice Requires="w"/><mc:Fallback>{fld(kind)}</mc:Fallback>'
        "</mc:AlternateContent>"
    ),
    "choice not understood": lambda kind: (
        f'<mc:AlternateContent><mc:Choice Requires="x">{fld(kind)}</mc:Choice>'
        "<mc:Fallback/></mc:AlternateContent>"
    ),
}


def field_text(*content: str) -> str:
    def fill(document: DocumentObject) -> None:
        add_xml(document, paragraph("".join(content)))

    return text_of(build(fill))


@pytest.mark.parametrize("place", NOT_SHOWN.values(), ids=NOT_SHOWN.keys())
def test_a_hidden_separate_still_ends_the_code(place: Callable[[str], str]) -> None:
    assert field_text(
        fld("begin"), code(CANARY), place("separate"), run("result"), fld("end"), run(" after")
    ) == ("result after")


@pytest.mark.parametrize("place", NOT_SHOWN.values(), ids=NOT_SHOWN.keys())
def test_a_hidden_end_still_closes_the_field(place: Callable[[str], str]) -> None:
    # A field nested in another field's code: if its hidden end were missed, the outer field's
    # separate would split the nested field instead, and the outer result would stay in the code.
    assert field_text(
        fld("begin"),
        code("IF"),
        fld("begin"),
        code("DATE"),
        fld("separate"),
        run(CANARY),
        place("end"),
        fld("separate"),
        run("result"),
        fld("end"),
        run(" after"),
    ) == ("result after")


@pytest.mark.parametrize("place", NOT_SHOWN.values(), ids=NOT_SHOWN.keys())
def test_a_hidden_begin_still_starts_the_code(place: Callable[[str], str]) -> None:
    # A field nested in the code of a field whose begin is hidden is still in the code.
    assert field_text(
        place("begin"),
        code("IF"),
        fld("begin"),
        code("DATE"),
        fld("separate"),
        run(CANARY),
        fld("end"),
        fld("separate"),
        run("result"),
        fld("end"),
        run(" after"),
    ) == ("result after")


@pytest.mark.parametrize("place", NOT_SHOWN.values(), ids=NOT_SHOWN.keys())
def test_nested_fields_with_every_marker_hidden_keep_only_the_outer_result(
    place: Callable[[str], str],
) -> None:
    assert field_text(
        run("before "),
        place("begin"),
        code("IF"),
        place("begin"),
        code("DATE"),
        place("separate"),
        run(CANARY),
        place("end"),
        place("separate"),
        run("outer"),
        place("begin"),
        code("PAGE"),
        place("separate"),
        run(" 7"),
        place("end"),
        place("end"),
        run(" after"),
    ) == ("before outer 7 after")


@pytest.mark.parametrize("place", NOT_PRESENT.values(), ids=NOT_PRESENT.keys())
def test_a_marker_that_is_not_in_the_accepted_document_does_not_count(
    place: Callable[[str], str],
) -> None:
    # Accepted, the field has no separate, so everything to its end is code.
    assert field_text(
        fld("begin"), code("REF a"), place("separate"), run(CANARY), fld("end"), run("after")
    ) == ("after")
    # Accepted, the outer field has no begin, so the nested field is an ordinary field.
    assert field_text(
        place("begin"),
        fld("begin"),
        code("DATE"),
        fld("separate"),
        run("nested"),
        fld("end"),
        fld("separate"),
        run(" result"),
        place("end"),
        run(" after"),
    ) == ("nested result after")


def test_a_marker_in_the_representation_that_is_read_counts() -> None:
    chosen = (
        f'<mc:AlternateContent><mc:Choice Requires="w">{fld("separate")}</mc:Choice>'
        f"<mc:Fallback>{fld('end')}</mc:Fallback></mc:AlternateContent>"
    )

    assert field_text(fld("begin"), code(CANARY), chosen, run("result"), fld("end")) == "result"


def test_a_text_box_anchored_in_a_hidden_run_is_left_out() -> None:
    hidden = (
        f"<w:r><w:rPr><w:vanish/></w:rPr><w:pict><w:txbxContent><w:p>{run(CANARY)}</w:p>"
        "</w:txbxContent></w:pict></w:r>"
    )

    assert field_text(run("Anchor"), hidden) == "Anchor"


def text_box_run(text: str) -> str:
    """A run anchoring a text box whose paragraph holds ``text``, a table and a nested box."""
    nested = (
        f"<w:p><w:r><w:pict><w:txbxContent><w:p>{run(text + ' nested')}</w:p></w:txbxContent>"
        "</w:pict></w:r></w:p>"
    )
    return (
        f"<w:r><w:pict><w:txbxContent><w:p>{run(text)}</w:p><w:tbl><w:tr>{cell(text + ' cell')}"
        f"</w:tr></w:tbl>{nested}</w:txbxContent></w:pict></w:r>"
    )


# Everything the walker gives content for, in a field's code: the INCLUDED_CONTAINERS cases (text
# and every character, through every container) and a text box.
EMITTING_CASES = [
    *((name, content, expected) for name, content, expected in INCLUDED_CONTAINERS),
    ("text box", text_box_run("boxed"), ""),
]


@pytest.mark.parametrize(
    ("content", "expected"),
    [pytest.param(content, expected, id=name) for name, content, expected in EMITTING_CASES],
)
def test_nothing_in_a_fields_code_is_given(content: str, expected: str) -> None:
    def in_code(inner: str) -> str:
        return f"{fld('begin')}{code('X')}{inner}{fld('separate')}{run('result')}{fld('end')}"

    assert field_text(in_code(content)) == "result"
    # In a field nested in the code, and in the code of a field nested in a result, too.
    nested = f"{fld('begin')}{code('IF')}{in_code(content)}{fld('separate')}{run('r')}{fld('end')}"
    assert field_text(nested) == "r"
    in_result = f"{fld('begin')}{code('A')}{fld('separate')}{in_code(content)}{fld('end')}"
    assert field_text(in_result) == "result"


def test_a_text_box_in_a_fields_result_is_read() -> None:
    content = f"{fld('begin')}{code('X')}{fld('separate')}{text_box_run('boxed')}{fld('end')}"

    assert field_text(run("Anchor"), content).split(PARAGRAPH_BREAK) == [
        "Anchor",
        "boxed",
        "boxed cell",
        "boxed nested",
    ]


def test_every_emitted_element_is_swept_in_a_fields_code() -> None:
    # The "characters" case writes w:t and every character; the "text box" case w:txbxContent.
    swept = {qn("w:t"), qn("w:txbxContent"), *CHARACTERS}

    assert swept == extractor.EMITTED
    assert {"characters", "text box"} <= {name for name, _, _ in EMITTING_CASES}


def test_every_excluded_element_is_decided_for_field_markers() -> None:
    # Each excluded element is either walked for its field markers (shown nowhere, but in the
    # document), or holds content accepting the changes removes, or has no runs to hold a marker.
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    w14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    removed = {f"{w}del", f"{w}moveFrom", f"{w14}conflictDel"}
    no_runs = {
        f"{w}{name}"
        for name in ("pPr", "rPr", "delText", "delInstrText", "instrText", "sdtPr", "sdtEndPr")
    } | {f"{w}sym"}

    assert {f"{w}rt", f"{m}oMath", f"{m}oMathPara"} == extractor.FIELDS_ONLY
    assert extractor.FIELDS_ONLY | removed | no_runs == EXCLUDED
    assert not extractor.FIELDS_ONLY & (removed | no_runs)
    assert not removed & no_runs


def test_japanese_text_with_ruby_reads_as_written() -> None:
    # A heading with furigana over each kanji: the heading's title is the base text alone.
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            + ruby("\u65e5\u672c", "\u306b\u307b\u3093")
            + run("\u306e")
            + ruby("\u6b74\u53f2", "\u308c\u304d\u3057")
            + "</w:p>",
        )

    page = extract(build(fill)).pages[0]

    assert page.text == "\u65e5\u672c\u306e\u6b74\u53f2"
    assert [(h.level, h.title) for h in page.headings] == [(1, page.text)]


# Text boxes


def text_box_paragraph(text: str, *, style: str | None = None) -> str:
    """A paragraph anchoring a text box, written the way Word does: a DrawingML shape in
    ``mc:Choice`` and the same box again as VML in ``mc:Fallback``."""
    style_xml = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    box = f"<w:txbxContent><w:p>{style_xml}<w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    return (
        f'<w:p {W} {MC} {WPS} {VML}><w:r><w:t xml:space="preserve">Anchor </w:t></w:r><w:r>'
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx>'
        f"{box}</wps:txbx></wps:wsp></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape>"
        f"<v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>"
        "</w:r><w:r><w:t>text.</w:t></w:r></w:p>"
    )


def test_a_text_box_is_read_once_after_its_anchor() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, text_box_paragraph("Boxed note"))
        document.add_paragraph("Next.")

    assert text_of(build(fill)) == (
        f"Anchor text.{PARAGRAPH_BREAK}Boxed note{PARAGRAPH_BREAK}Next."
    )


def test_a_vml_only_text_box_is_read() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {VML}><w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r>"
            "<w:t>Old box</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>"
            "</w:r></w:p>",
        )

    assert text_of(build(fill)) == "Old box"


# The namespaces Word writes in ``mc:Choice Requires``, each decided: (prefix, namespace).
WORD_NAMESPACES = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
    "wpi": "http://schemas.microsoft.com/office/word/2010/wordprocessingInk",
    "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "w15": "http://schemas.microsoft.com/office/word/2012/wordml",
    "w16se": "http://schemas.microsoft.com/office/word/2015/wordml/symex",
    "cx1": "http://schemas.microsoft.com/office/drawing/2015/9/8/chartex",
}
DECLARED = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in WORD_NAMESPACES.items())
PICTURE = (
    '<w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
    'wordprocessingDrawing"><wp:docPr id="1" name="Picture"/></wp:inline></w:drawing>'
)


def box(text: str) -> str:
    return f"<w:txbxContent><w:p>{run(text)}</w:p></w:txbxContent>"


def shape(text: str) -> str:
    return f"<wps:wsp><wps:spPr/><wps:txbx>{box(text)}</wps:txbx><wps:bodyPr/></wps:wsp>"


def graphic(prefix: str, content: str) -> str:
    return (
        '<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
        'wordprocessingDrawing"><a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        f'2006/main"><a:graphicData uri="{WORD_NAMESPACES[prefix]}">{content}'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
    )


# For each understood namespace, content as Word writes it with text inside, in a choice whose
# fallback is only a picture: the text is found only by reading the choice. (the choice's content,
# the paragraphs it gives after the anchor's own text)
TEXT_BEARING = {
    "wps": (graphic("wps", shape("in a shape")), ["Anchor", "in a shape"]),
    "wpg": (
        graphic("wpg", f"<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr/>{shape('in a group')}</wpg:wgp>"),
        ["Anchor", "in a group"],
    ),
    "wpc": (
        graphic(
            "wpc",
            "<wpc:wpc><wpc:bg/><wpc:whole/>"
            f"{shape('in a canvas')}"
            f"<wpg:wgp><wpg:grpSpPr/>{shape('in a group in a canvas')}</wpg:wgp></wpc:wpc>",
        ),
        ["Anchor", "in a canvas", "in a group in a canvas"],
    ),
    "w14": (
        f'<w14:conflictIns w:id="1" w:author="A">{run(" agreed")}</w14:conflictIns>',
        ["Anchor agreed"],
    ),
}


@pytest.mark.parametrize("prefix", TEXT_BEARING)
def test_text_in_every_understood_namespace_is_read(prefix: str) -> None:
    content, expected = TEXT_BEARING[prefix]

    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {MC} {DECLARED}>{run('Anchor')}<mc:AlternateContent>"
            f'<mc:Choice Requires="{prefix}">{content}</mc:Choice>'
            f"<mc:Fallback><w:r>{PICTURE}</w:r></mc:Fallback></mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)).split(PARAGRAPH_BREAK) == expected


# Each namespace that is not understood, with what Word writes in the fallback for it.
NOT_UNDERSTOOD = {
    "wpi": "",  # ink: a picture, no text
    "wp14": "",  # relative position or size: the absolute one, no text
    "a14": "",  # DrawingML extension: no text
    "w15": "same text",
    "w16se": "\U0001f600",  # an emoji: w16se:symEx names it, the fallback writes it
    "cx1": "",  # chartex: a picture; the chart's text is in its own part
}


@pytest.mark.parametrize("prefix", NOT_UNDERSTOOD)
def test_a_choice_requiring_a_namespace_not_understood_falls_back(prefix: str) -> None:
    fallback = run(NOT_UNDERSTOOD[prefix]) if NOT_UNDERSTOOD[prefix] else PICTURE

    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {MC} {DECLARED}>{run('[')}<mc:AlternateContent>"
            f'<mc:Choice Requires="{prefix}"><w:r><w:drawing>{shape(CANARY)}</w:drawing></w:r>'
            f"{run(CANARY)}</mc:Choice><mc:Fallback>{fallback}</mc:Fallback>"
            f"</mc:AlternateContent>{run(']')}</w:p>",
        )

    assert text_of(build(fill)) == f"[{NOT_UNDERSTOOD[prefix]}]"


def test_every_namespace_word_requires_is_decided() -> None:
    assert set(TEXT_BEARING) | set(NOT_UNDERSTOOD) == set(WORD_NAMESPACES)
    assert not set(TEXT_BEARING) & set(NOT_UNDERSTOOD)
    assert {WORD_NAMESPACES[prefix] for prefix in TEXT_BEARING} == extractor.UNDERSTOOD - {
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }


# mc:Ignorable, mc:ProcessContent and mc:MustUnderstand (ECMA-376 Part 3), checked against .NET's
# XmlCompatibilityReader, the reference reader: an element's own attributes apply to itself.

MC_URI = "http://schemas.openxmlformats.org/markup-compatibility/2006"
IGNORE_X = 'mc:Ignorable="x"'


def add_blocks(document: DocumentObject, body_xml: str) -> None:
    """Append the children of a ``w:body`` written as XML (so they can share its namespace
    declarations) to the document's body, before the section properties."""
    body = document.element.body
    for block in list(parse_xml(body_xml)):
        body.insert(len(body) - 1, block)


LATER = f"<x:later>{run(CANARY)}</x:later>"


def nested_in_a_text_box_in_a_cell(levels: dict[str, str], wrapper: str = LATER) -> str:
    """A table whose cell holds a paragraph anchoring a text box whose paragraph holds
    ``wrapper``. ``levels`` gives extra attributes for any of: table, row, cell, paragraph, run,
    box, inner (the text box's paragraph)."""

    def at(level: str) -> str:
        return " " + levels[level] if level in levels else ""

    return (
        f"<w:tbl {W} {MC} {WPS} {FUTURE}{at('table')}><w:tr{at('row')}><w:tc{at('cell')}>"
        f"<w:p{at('paragraph')}>{run('Anchor')}<w:r{at('run')}><w:drawing><wps:wsp><wps:txbx>"
        f"<w:txbxContent{at('box')}><w:p{at('inner')}>{wrapper}{run('boxed')}</w:p>"
        "</w:txbxContent></wps:txbx></wps:wsp></w:drawing></w:r></w:p></w:tc></w:tr></w:tbl>"
    )


@pytest.mark.parametrize(
    "level", ["table", "row", "cell", "paragraph", "run", "box", "inner", "wrapper itself"]
)
def test_an_ignorable_element_is_dropped_whichever_ancestor_declares_it(level: str) -> None:
    wrapper = f"<x:later {IGNORE_X}>{run(CANARY)}</x:later>" if level == "wrapper itself" else LATER

    def fill(document: DocumentObject) -> None:
        add_xml(document, nested_in_a_text_box_in_a_cell({level: IGNORE_X}, wrapper))

    assert text_of(build(fill)) == "Anchor boxed"


def test_without_mc_ignorable_an_unknown_wrapper_is_read_through() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, nested_in_a_text_box_in_a_cell({}, f"<x:later>{run('kept ')}</x:later>"))

    assert text_of(build(fill)) == "Anchor kept boxed"


def test_ignorable_declared_on_the_body_and_on_the_document_root_is_inherited() -> None:
    # The template's root declares ``mc:Ignorable="w14 wp14"``, as Word does; the body here
    # declares the ink namespace, whose prefix the root binds.
    wpi = 'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk"'
    wp14 = 'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"'

    def fill(document: DocumentObject) -> None:
        body = document.element.body
        assert "wp14" in document.element.get(f"{{{MC_URI}}}Ignorable", "").split()
        body.set(f"{{{MC_URI}}}Ignorable", "wpi")
        add_xml(document, paragraph(f"{run('a')}<wpi:x {wpi}>{run(CANARY)}</wpi:x>{run('b')}"))
        add_xml(document, paragraph(f"{run('c')}<wp14:x {wp14}>{run(CANARY)}</wp14:x>"))

    assert text_of(build(fill)) == f"ab{PARAGRAPH_BREAK}c"


@pytest.mark.parametrize(
    "prefix",
    [
        pytest.param("w14", id="w14 (Word declares it ignorable in every document)"),
        pytest.param("w", id="WordprocessingML"),
    ],
)
def test_an_understood_namespace_declared_ignorable_is_never_ignored(prefix: str) -> None:
    content = {
        "w14": f'<w14:conflictIns w:id="1" w:author="A">{run("agreed")}</w14:conflictIns>',
        "w": run("plain"),
    }[prefix]

    def fill(document: DocumentObject) -> None:
        add_xml(document, f'<w:p {W} {MC} {W14} mc:Ignorable="{prefix}">{content}</w:p>')

    assert text_of(build(fill)) == {"w14": "agreed", "w": "plain"}[prefix]


def test_a_text_box_shape_declared_ignorable_is_still_read() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            nested_in_a_text_box_in_a_cell({"paragraph": 'mc:Ignorable="wps"'}, run("in box ")),
        )

    assert text_of(build(fill)) == "Anchor in box boxed"


def test_a_prefix_is_resolved_where_mc_ignorable_is_written() -> None:
    # ``x`` means urn:a where the attribute is; deeper, ``x`` is bound to urn:b, which is not
    # ignorable, and urn:a is written with another prefix, which still is.
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f'<w:p {W} {MC} xmlns:x="urn:a" mc:Ignorable="x">{run("a")}'
            f'<z:dropped xmlns:z="urn:a">{run(CANARY)}</z:dropped>'
            f'<x:kept xmlns:x="urn:b">{run("b")}</x:kept></w:p>',
        )

    assert text_of(build(fill)) == "ab"


@pytest.mark.parametrize(
    ("process", "expected"),
    [
        pytest.param('mc:ProcessContent="x:later"', "[in]", id="named"),
        pytest.param('mc:ProcessContent="x:*"', "[in]", id="wildcard"),
        pytest.param('mc:ProcessContent="x:other"', "[]", id="another name"),
        pytest.param("", "[]", id="none"),
    ],
)
def test_process_content_keeps_the_content_of_an_ignored_element(
    process: str, expected: str
) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            paragraph(f"{run('[')}<x:later {IGNORE_X} {process}>{run('in')}</x:later>{run(']')}"),
        )

    assert text_of(build(fill)) == expected


def test_redeclaring_a_namespace_ignorable_keeps_the_inherited_process_content() -> None:
    # As in XmlCompatibilityReader: mc:Ignorable adds to the scope's ignorable namespaces and does
    # not touch ProcessContent, which is looked up from the nearest scope that has an entry for the
    # namespace. So an inner mc:Ignorable="x" alone keeps the table's x:* processing.
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            nested_in_a_text_box_in_a_cell(
                {"table": 'mc:Ignorable="x" mc:ProcessContent="x:*"', "inner": IGNORE_X},
                f"<x:a>{run('a ')}</x:a>",
            ),
        )

    assert text_of(build(fill)) == "Anchor a boxed"


def test_process_content_is_inherited_and_the_nearest_declaration_for_a_namespace_wins() -> None:
    # The table processes every x element; the inner paragraph names only x:a for x, so x:b
    # there is dropped with its content, while x:b in the outer paragraph is processed.
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            nested_in_a_text_box_in_a_cell(
                {
                    "table": 'mc:Ignorable="x" mc:ProcessContent="x:*"',
                    "inner": 'mc:Ignorable="x" mc:ProcessContent="x:a"',
                },
                f"<x:a>{run('a ')}</x:a><x:b>{run(CANARY)}</x:b>",
            ),
        )

    assert text_of(build(fill)) == "Anchor a boxed"


@pytest.mark.parametrize("wrapped", ["paragraph", "row", "cell"])
@pytest.mark.parametrize("processed", [True, False], ids=["processed", "not processed"])
def test_an_ignorable_wrapper_around_a_block_is_read_through_only_when_processed(
    wrapped: str, processed: bool
) -> None:
    attributes = f'{IGNORE_X} mc:ProcessContent="x:wrap"' if processed else IGNORE_X
    text = "wrapped" if processed else CANARY

    def wrap(xml: str) -> str:
        return f"<x:wrap {attributes}>{xml}</x:wrap>"

    cell_xml = cell(text)
    row_xml = f"<w:tr>{cell('first')}{cell_xml}</w:tr>"
    blocks = {
        "paragraph": wrap(f"<w:p>{run(text)}</w:p>"),
        "row": f"<w:tbl><w:tr>{cell('first')}</w:tr>{wrap(row_xml)}</w:tbl>",
        "cell": f"<w:tbl><w:tr>{cell('first')}{wrap(cell_xml)}</w:tr></w:tbl>",
    }[wrapped]

    def fill(document: DocumentObject) -> None:
        add_blocks(document, f"<w:body {W} {MC} {FUTURE}>{blocks}</w:body>")

    expected = {
        ("paragraph", True): "wrapped",
        ("paragraph", False): "",
        ("row", True): "first\nfirst\twrapped",
        ("row", False): "first",
        ("cell", True): "first\twrapped",
        ("cell", False): "first",
    }[wrapped, processed]
    assert text_of(build(fill)) == expected


def test_a_heading_style_inside_an_ignorable_wrapper_counts_only_when_processed() -> None:
    def styles_with(attributes: str) -> Callable[[DocumentObject], None]:
        def fill(document: DocumentObject) -> None:
            document.styles.element.append(
                parse_xml(
                    f"<x:wrap {W} {MC} {FUTURE} {attributes}>"
                    '<w:style w:type="paragraph" w:styleId="Wrapped"><w:name w:val="heading 3"/>'
                    "</w:style></x:wrap>"
                )
            )
            add_xml(
                document,
                f'<w:p {W}><w:pPr><w:pStyle w:val="Wrapped"/></w:pPr>{run("Section")}</w:p>',
            )

        return fill

    processed = extract(build(styles_with(f'{IGNORE_X} mc:ProcessContent="x:wrap"'))).pages[0]
    ignored = extract(build(styles_with(IGNORE_X))).pages[0]
    read_through = extract(build(styles_with(""))).pages[0]

    assert [(h.level, h.title) for h in processed.headings] == [(3, "Section")]
    assert ignored.headings == ()
    assert read_through.headings == ()  # the style reader looks only at the styles' own children


def test_content_of_an_ignored_element_is_not_examined() -> None:
    # Imported content inside an ignored element is gone, so the file is not refused for it.
    def fill(document: DocumentObject) -> None:
        add_blocks(
            document,
            f'<w:body {W} {MC} {FUTURE} {R}><x:wrap {IGNORE_X}><w:altChunk r:id="rIdImported"/>'
            '<w:p mc:MustUnderstand="x"/></x:wrap></w:body>',
        )

    assert text_of(build(fill)) == ""


@pytest.mark.parametrize(
    "attributes",
    [
        pytest.param('mc:MustUnderstand="w"', id="WordprocessingML"),
        pytest.param('mc:MustUnderstand="w w14 wps wpg wpc"', id="every understood namespace"),
    ],
)
def test_must_understand_naming_only_understood_namespaces_is_read(attributes: str) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, f"<w:p {W} {MC} {DECLARED} {attributes}>{run('ok')}</w:p>")

    assert text_of(build(fill)) == "ok"


@pytest.mark.parametrize(
    "markup",
    [
        pytest.param(f'<w:p {W} {MC} {FUTURE} mc:MustUnderstand="x"/>', id="must understand"),
        pytest.param(
            f'<w:p {W} {MC} {FUTURE} mc:MustUnderstand="w x"/>', id="must understand, one of two"
        ),
        pytest.param(
            f'<w:p {W} {MC} {W15} mc:MustUnderstand="w15"/>', id="must understand a Word extension"
        ),
        pytest.param(
            f'<w:p {W} {MC} mc:MustUnderstand="nowhere"/>', id="must understand undeclared"
        ),
        pytest.param(f'<w:p {W} {MC} mc:Ignorable="nowhere"/>', id="ignorable undeclared"),
        pytest.param(
            f'<w:p {W} {MC} {FUTURE} {IGNORE_X} mc:ProcessContent="y:a"/>',
            id="process content undeclared",
        ),
        pytest.param(
            f'<w:p {W} {MC} {FUTURE} mc:ProcessContent="x:a"/>', id="process content not ignorable"
        ),
        pytest.param(
            f'<w:tbl {W} {MC} {FUTURE} {IGNORE_X}><w:tr mc:ProcessContent="x:a"/></w:tbl>',
            id="process content ignorable only on an ancestor",
        ),
        *(
            pytest.param(
                f'<w:p {W} {MC} {FUTURE} {IGNORE_X} mc:ProcessContent="{entry}"/>',
                id=f"process content {entry!r}",
            )
            for entry in ["x", "x:", ":a", "x:a:b", "x:1a", "x:a b:"]
        ),
        pytest.param(
            f"<w:p {W} {MC}><mc:Choice Requires='w'>{run(CANARY)}</mc:Choice></w:p>",
            id="choice outside alternate content",
        ),
        pytest.param(
            f"<w:p {W} {MC}><mc:Fallback>{run(CANARY)}</mc:Fallback></w:p>",
            id="fallback outside alternate content",
        ),
        pytest.param(f"<w:p {W} {MC}><mc:Unknown/></w:p>", id="unknown compatibility element"),
        pytest.param(
            f"<w:p {W} {MC} {FUTURE}><w:r><mc:AlternateContent><mc:Choice Requires='w'>"
            f'<w:t mc:MustUnderstand="x">{CANARY}</w:t></mc:Choice></mc:AlternateContent></w:r>'
            "</w:p>",
            id="must understand inside the chosen representation",
        ),
    ],
)
def test_compatibility_markup_that_is_an_error_refuses_the_file(markup: str) -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, markup)

    with pytest.raises(CompatibilityError) as caught:
        extract(build(fill))
    assert (
        str(caught.value) == "The DOCX uses markup compatibility rules this extractor cannot apply."
    )
    assert CANARY not in str(caught.value)


def test_must_understand_in_a_representation_that_is_not_read_is_not_examined() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(
            document,
            f"<w:p {W} {MC} {FUTURE}><mc:AlternateContent><mc:Choice Requires='w'>{run('read')}"
            f'</mc:Choice><mc:Fallback><w:r mc:MustUnderstand="x"/></mc:Fallback>'
            "</mc:AlternateContent></w:p>",
        )

    assert text_of(build(fill)) == "read"


def test_comments_and_processing_instructions_are_passed_over() -> None:
    def fill(document: DocumentObject) -> None:
        add_xml(document, paragraph(f"{run('a')}<!-- {CANARY} --><?pi {CANARY}?>{run('b')}"))

    assert text_of(build(fill)) == "ab"


def test_a_body_inside_an_ignored_element_leaves_no_document() -> None:
    document_xml = (
        f"<w:document {W} {MC} {FUTURE} {IGNORE_X}><x:wrap><w:body><w:p>{run(CANARY)}</w:p>"
        "</w:body></x:wrap></w:document>"
    ).encode()

    with pytest.raises(NotDocxError):
        extract(replace_member(build(), "word/document.xml", document_xml))


def test_a_must_understand_error_in_the_styles_part_refuses_the_file() -> None:
    def fill(document: DocumentObject) -> None:
        document.styles.element.append(parse_xml(f'<x:s {W} {MC} {FUTURE} mc:MustUnderstand="x"/>'))

    with pytest.raises(CompatibilityError):
        extract(build(fill))


# Refused input


def test_errors_are_extraction_errors_with_fixed_text() -> None:
    assert issubclass(NotDocxError, ExtractionError)
    assert str(NotDocxError(Reason.UNSAFE_ARCHIVE)) == (
        "The file is not a readable DOCX (unsafe_archive)."
    )
    assert str(NotDocxError()) == "The file is not a readable DOCX (unreadable_part)."


@pytest.mark.parametrize(
    ("data", "reason"),
    [
        (b"", Reason.EMPTY),
        (OLE_SIGNATURE + b"\x00" * 504, Reason.LEGACY_OR_ENCRYPTED_OFFICE),
        (b"PK\x03\x04" + bytes(range(200)), Reason.UNREADABLE_ARCHIVE),
        (f"{CANARY} plain text".encode(), Reason.CONTENT_MISMATCH),
        (b"%PDF-1.7\n%%EOF\n", Reason.CONTENT_MISMATCH),
    ],
)
def test_bytes_that_are_not_a_docx(data: bytes, reason: Reason, library_never_called: None) -> None:
    with pytest.raises(NotDocxError) as raised:
        extract(data)

    assert raised.value.reason is reason
    assert CANARY not in str(raised.value)


def test_a_zip_that_is_not_a_docx(library_never_called: None) -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("hello.txt", CANARY)

    with pytest.raises(NotDocxError) as raised:
        extract(out.getvalue())

    assert raised.value.reason is Reason.CONTENT_MISMATCH


def test_a_docx_whose_xml_is_broken_gives_fixed_text_without_the_library_error() -> None:
    data = replace_member(build(), "word/document.xml", f"<w:document>{CANARY}".encode())

    with pytest.raises(NotDocxError) as raised:
        extract(data)

    assert raised.value.reason is None
    assert CANARY not in str(raised.value)
    assert raised.value.__cause__ is None
    assert raised.value.__suppress_context__


def test_a_package_whose_main_part_is_not_a_word_document_is_refused() -> None:
    data = build()
    types = dict(members_of(data))["[Content_Types].xml"]
    macro = types.replace(
        b"wordprocessingml.document.main+xml", b"ms-word.document.macroEnabled.main+xml"
    )
    assert macro != types

    with pytest.raises(NotDocxError):
        extract(replace_member(data, "[Content_Types].xml", macro))


def test_a_strict_open_xml_document_is_refused_not_read_as_empty() -> None:
    data = build(lambda d: d.add_paragraph("Strict text."))
    document_xml = dict(members_of(data))["word/document.xml"]
    strict = document_xml.replace(
        b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        b"http://purl.oclc.org/ooxml/wordprocessingml/main",
    )
    assert strict != document_xml

    with pytest.raises(NotDocxError) as raised:
        extract(replace_member(data, "word/document.xml", strict))

    assert raised.value.reason is None


def test_a_member_whose_crc_does_not_match_is_unreadable() -> None:
    members = [
        (name, deflate(content), zipfile.ZIP_DEFLATED, zlib.crc32(content) ^ 1, len(content))
        for name, content in members_of(build())
    ]

    with pytest.raises(NotDocxError) as raised:
        extract(raw_zip(members))

    assert raised.value.reason is Reason.UNREADABLE_ARCHIVE


def test_folder_entries_in_the_zip_are_skipped() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/", b"")
        for name, content in members_of(build(lambda d: d.add_paragraph("With folders."))):
            archive.writestr(name, content)

    assert text_of(out.getvalue()) == "With folders."


def test_a_member_is_never_read_past_its_claim_even_if_zipfile_did() -> None:
    class Archive:
        def open(self, _: zipfile.ZipInfo) -> io.BytesIO:
            return io.BytesIO(b"x" * 10)

    info = zipfile.ZipInfo("word/document.xml")
    info.file_size = 4

    with pytest.raises(NotDocxError) as raised:
        extractor._read_member(Archive(), info)  # type: ignore[arg-type]

    assert raised.value.reason is Reason.UNSAFE_ARCHIVE


def test_running_out_of_memory_is_not_reported_as_a_bad_document(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def exhausted(*_: object) -> NoReturn:
        raise MemoryError

    monkeypatch.setattr(extractor, "Document", exhausted)

    with pytest.raises(MemoryError):
        extract(build())


def test_check_docx_agrees_with_the_import_check() -> None:
    assert check_docx(build()) is None
    assert check_docx(b"x" * 11, Limits(max_bytes=10)) is Reason.TOO_LARGE


# Zip bombs


def test_a_compression_bomb_is_refused_before_python_docx_opens_it(
    library_never_called: None,
) -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in members_of(build()):
            archive.writestr(name, content)
        archive.writestr("word/media/zeros.bin", b"\x00" * 20_000_000)
    assert len(out.getvalue()) < 100_000

    with pytest.raises(NotDocxError) as raised:
        extract(out.getvalue())

    assert raised.value.reason is Reason.UNSAFE_ARCHIVE


def test_a_directory_over_the_limits_is_refused_before_zipfile_opens_it(
    library_never_called: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    def refuse(*_: object, **__: object) -> NoReturn:
        raise AssertionError("zipfile was opened")

    monkeypatch.setattr("vault.ingest.validate.zipfile.ZipFile", refuse)

    with pytest.raises(NotDocxError) as raised:
        extract(build(), Limits(max_zip_entries=3))

    assert raised.value.reason is Reason.UNSAFE_ARCHIVE


def test_a_member_longer_than_its_claim_is_never_decompressed_past_the_claim() -> None:
    """zipfile's ``read`` decompresses a whole deflate stream before cutting it to the claimed size,
    so a stream far longer than its claim would fill memory inside python-docx. The extractor reads
    members in pieces and hands python-docx a copy that holds only the claimed bytes."""

    def fill(document: DocumentObject) -> None:
        document.add_paragraph("The real text.")

    padding = 64 * 1024 * 1024
    members = []
    for name, content in members_of(build(fill)):
        stream = content + b" " * padding if name == "word/document.xml" else content
        # The CRC and size of the claimed prefix only, so the claim checks out where it stops.
        members.append(
            (name, deflate(stream), zipfile.ZIP_DEFLATED, zlib.crc32(content), len(content))
        )
    data = raw_zip(members)
    assert check_docx(data) is None  # the claims are small, so validation passes

    # Control: reading the member the way python-docx does allocates the whole stream.
    tracemalloc.start()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        archive.read("word/document.xml")
    _, unbounded = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert unbounded > padding

    tracemalloc.start()
    try:
        result = extract(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert result.pages[0].text == "The real text."
    assert peak < padding // 4
